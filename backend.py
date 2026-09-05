from __future__ import annotations
import io, os, re, shutil, tempfile, unicodedata, subprocess, platform
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy import create_engine, Column, Integer, String, Text, DateTime, ForeignKey
from sqlalchemy.orm import declarative_base, sessionmaker
from sqlalchemy.exc import SQLAlchemyError
import hashlib, hmac, secrets, json
from datetime import datetime
from pydantic import BaseModel

import fitz
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.units import cm

try:
    import pytesseract
except Exception:
    pytesseract = None

ROOT = Path(__file__).resolve().parent
app = FastAPI(title='APM Contratos Web', version='0.6')
app.add_middleware(SessionMiddleware, secret_key=os.getenv('APP_SECRET','dev-change-me'), https_only=os.getenv('COOKIE_SECURE','false').lower()=='true', same_site='lax')


def find_tesseract():
    found = shutil.which('tesseract')
    if found:
        return found
    if os.name == 'nt':
        candidates = [
            Path(os.environ.get('ProgramFiles', r'C:\\Program Files')) / 'Tesseract-OCR' / 'tesseract.exe',
            Path(os.environ.get('LOCALAPPDATA', '')) / 'Programs' / 'Tesseract-OCR' / 'tesseract.exe',
            Path(r'C:\\Program Files (x86)\Tesseract-OCR\tesseract.exe'),
        ]
        for c in candidates:
            if c and c.exists():
                return str(c)
    return None

def configure_tesseract():
    exe = find_tesseract()
    if exe and pytesseract:
        try:
            pytesseract.pytesseract.tesseract_cmd = exe
        except Exception:
            pass
    return exe

def tesseract_available():
    return bool(pytesseract and configure_tesseract())

@app.get('/api/health')
def health():
    exe = configure_tesseract()
    return {'ok': True, 'version': '0.6', 'mode':'web-production', 'ocr': bool(exe and pytesseract), 'tesseract_path': exe}

@app.get('/api/ocr-status')
def ocr_status():
    exe = configure_tesseract()
    return {
        'installed': bool(exe),
        'python_bridge': bool(pytesseract),
        'ready': bool(exe and pytesseract),
        'path': exe,
        'platform': platform.system(),
        'can_auto_install': os.name == 'nt' and bool(shutil.which('winget'))
    }

@app.post('/api/install-tesseract')
def install_tesseract():
    if os.name != 'nt':
        return JSONResponse({'ok': False, 'message': 'A instalação automática desta versão é destinada ao Windows.'}, status_code=400)
    winget = shutil.which('winget')
    if not winget:
        return JSONResponse({'ok': False, 'message': 'Winget não foi encontrado. Instale o Tesseract manualmente e reinicie o APM.'}, status_code=400)
    commands = [
        [winget, 'install', '-e', '--id', 'UB-Mannheim.TesseractOCR', '--accept-package-agreements', '--accept-source-agreements', '--silent'],
        [winget, 'install', '-e', '--id', 'tesseract-ocr.tesseract', '--accept-package-agreements', '--accept-source-agreements', '--silent'],
    ]
    last = ''
    for cmd in commands:
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
            last = (proc.stdout or '') + '\n' + (proc.stderr or '')
            if proc.returncode == 0 or find_tesseract():
                exe = configure_tesseract()
                return {'ok': bool(exe), 'message': 'Tesseract instalado e detectado.' if exe else 'Instalação concluída. Reinicie o APM para atualizar o PATH.', 'path': exe}
        except Exception as e:
            last = str(e)
    return JSONResponse({'ok': False, 'message': 'Não foi possível instalar automaticamente o Tesseract.', 'detail': last[-1200:]}, status_code=500)


def normalize(s: str) -> str:
    return re.sub(r'\s+', ' ', s or '').strip()


def classify(name: str, text: str, declared: str) -> str:
    if declared and declared != 'auto': return declared
    hay = (name + '\n' + text[:5000]).lower()
    checks = [
        ('Certidão de casamento', ['certidão de casamento','certidao de casamento']),
        ('Certidão de nascimento', ['certidão de nascimento','certidao de nascimento']),
        ('CIN', ['carteira de identidade nacional','documento de identificação']),
        ('CNH', ['carteira nacional de habilitação','secretaria nacional de trânsito','senatran']),
        ('RG', ['carteira de identidade','registro geral']),
        ('CPF', ['cadastro de pessoas físicas','cadastro de pessoas fisicas']),
        ('BCI/IPTU', ['boletim de cadastro imobiliário','inscrição imobiliária','valor venal']),
        ('Matrícula', ['registro de imóveis','registro de imoveis','matrícula','matricula','livro nº 2','livro n.º 2']),
        ('Comprovante de endereço', ['amazonas energia','manaus ambiental','unidade consumidora','conta de energia','conta de água']),
    ]
    for typ, words in checks:
        if any(w in hay for w in words): return typ
    return 'Documento'


def pdf_text(data: bytes):
    doc = fitz.open(stream=data, filetype='pdf')
    texts=[]
    for page in doc:
        texts.append(page.get_text('text') or '')
    return '\n'.join(texts), len(doc), doc


def ocr_doc(doc: fitz.Document, max_pages=12):
    if not tesseract_available(): return ''
    out=[]
    # Documentos pessoais costumam ocupar uma única página e se beneficiam de renderização maior.
    scale = 3.3 if len(doc) <= 2 else 2.1
    for i in range(min(len(doc), max_pages)):
        pix=doc[i].get_pixmap(matrix=fitz.Matrix(scale,scale), alpha=False)
        img=Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB')
        try:
            txt=pytesseract.image_to_string(img, lang='por')
            if not txt.strip(): txt=pytesseract.image_to_string(img)
            out.append(txt)
        except Exception:
            try: out.append(pytesseract.image_to_string(img))
            except Exception: out.append('')
    return '\n'.join(out)


def image_ocr(data: bytes):
    if not tesseract_available(): return ''
    img=Image.open(io.BytesIO(data)).convert('RGB')
    try: return pytesseract.image_to_string(img, lang='por')
    except Exception: return pytesseract.image_to_string(img)


def first_match(patterns, text, flags=re.I|re.M):
    for pat in patterns:
        m=re.search(pat,text,flags)
        if m: return normalize(m.group(1))
    return None


def add(fields, key, label, value, source):
    if value and not any(x['key']==key for x in fields):
        fields.append({'key':key,'label':label,'value':value,'source':source})


def extract_fields(typ: str, text: str, source: str):
    t=text.replace('\r','\n')
    fields=[]; pending=[]
    cpf = first_match([r'CPF(?:/MF)?\s*(?:n[ºo°.]?\s*)?[:\-]?\s*(\d{3}[.\s]?\d{3}[.\s]?\d{3}[-\s]?\d{2})', r'\b(\d{3}\.\d{3}\.\d{3}-\d{2})\b'], t)
    rg = first_match([r'(?:RG|REGISTRO GERAL|DOC(?:UMENTO)? DE IDENTIDADE)\s*(?:N[ºO°.]?\s*)?[:\-]?\s*([0-9.\-]+\s*(?:SSP|SESP|PC|IICC|SSDS)?/?[A-Z]{0,2})'],t)
    nasc=first_match([r'(?:DATA DE NASCIMENTO|NASCIMENTO)\s*[:\-]?\s*(\d{2}/\d{2}/\d{4})', r'\b(\d{2}/\d{2}/\d{4})\b'],t)
    if typ in ('CNH','RG','CIN','CPF'):
        name=first_match([r'(?:NOME(?: E SOBRENOME)?|NOME / NAME)\s*[:\-]?\s*\n?\s*([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][A-ZÁÀÂÃÉÊÍÓÔÕÚÇ ]{5,})',r'\n([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ]{2,}(?: [A-ZÁÀÂÃÉÊÍÓÔÕÚÇ]{2,}){2,})\n'],t)
        add(fields,'nome','Nome',name,source); add(fields,'cpf','CPF',cpf,source); add(fields,'rg','RG',rg,source); add(fields,'nascimento','Nascimento',nasc,source)
        cnh=first_match([r'(?:N[ºO°.]?\s*REGISTRO|REGISTRO)\s*[:\-]?\s*(\d{9,12})'],t)
        if typ=='CNH': add(fields,'cnh','Registro CNH',cnh,source)
        nat=first_match([r'NACIONALIDADE\s*[:\-]?\s*([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ]+)'],t)
        add(fields,'nacionalidade','Nacionalidade',nat,source)
        if not cpf: pending.append('CPF não localizado com segurança.')
        if not name: pending.append('Nome não localizado com segurança.')
        pending += ['Estado civil deve ser confirmado em certidão ou informado manualmente.','Profissão deve ser confirmada em documento próprio ou informada manualmente.']
    elif typ=='Certidão de casamento':
        regime=first_match([r'REGIME DE BENS[^\n]*\n\s*([^\n]{4,80})',r'(COMUNH[AÃ]O PARCIAL DE BENS|COMUNH[AÃ]O UNIVERSAL DE BENS|SEPARA[CÇ][AÃ]O[^\n]{0,40})'],t)
        names=re.findall(r'\b([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ]{2,}(?:\s+[A-ZÁÀÂÃÉÊÍÓÔÕÚÇ]{2,}){2,})\b',t)
        # remove common headings
        names=[n for n in names if not any(h in n for h in ['REPÚBLICA FEDERATIVA','REGISTRO CIVIL','CERTIDÃO DE CASAMENTO','COMUNHÃO PARCIAL'])]
        if names: add(fields,'conjuge1','Cônjuge 1',names[0],source)
        if len(names)>1: add(fields,'conjuge2','Cônjuge 2',names[1],source)
        add(fields,'regime','Regime de bens',regime,source); add(fields,'estado_civil','Estado civil','Casado(a)',source)
        data=first_match([r'(?:DATA DO REGISTRO|DATA DA CELEBRA[CÇ][AÃ]O)[^\n]*\n?\s*([^\n]{5,60})'],t)
        add(fields,'data','Data do casamento/registro',data,source)
        if not regime: pending.append('Regime de bens não localizado com segurança.')
    elif typ=='Certidão de nascimento':
        name=first_match([r'NOME\s*[:\-]?\s*\n?\s*([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][A-ZÁÀÂÃÉÊÍÓÔÕÚÇ ]{5,})'],t)
        add(fields,'nome','Nome',name,source); add(fields,'nascimento','Nascimento',nasc,source)
        fil=first_match([r'FILIA[CÇ][AÃ]O\s*[:\-]?\s*\n?\s*([^\n]{8,160})'],t)
        add(fields,'filiacao','Filiação',fil,source)
    elif typ=='Comprovante de endereço':
        cep=first_match([r'CEP\s*[:\-]?\s*(\d{2}\.?\d{3}-?\d{3})',r'\b(\d{5}-\d{3})\b'],t)
        endereco=first_match([r'\n((?:AV|AVENIDA|RUA|TRAVESSA|ALAMEDA|ESTRADA)\.?[^\n]{5,100})'],t)
        titular=first_match([r'\n([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ]{2,}(?: [A-ZÁÀÂÃÉÊÍÓÔÕÚÇ]{2,}){2,})\n'],t)
        add(fields,'titular','Titular',titular,source); add(fields,'endereco','Endereço',endereco,source); add(fields,'cep','CEP',cep,source)
        cidadeuf=first_match([r'\b([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ ]{3,})\s*[-/]\s*([A-Z]{2})\b'],t)
        if cidadeuf: add(fields,'cidade_uf','Cidade/UF',cidadeuf,source)
        if not endereco: pending.append('Endereço não localizado com segurança; confira a imagem do comprovante.')
    elif typ in ('Matrícula','BCI/IPTU'):
        mat=first_match([r'MATR[ÍI]CULA\s*(?:N[ºO°.]?\s*)?[:\-]?\s*([0-9.\-]{3,})',r'Matrícula\s*\n\s*([0-9.\-]{3,})'],t)
        add(fields,'matricula','Matrícula',mat,source)
        cart=first_match([r'((?:\d+[º°]?\s*)?OF[IÍ]CIO DE REGISTRO DE IM[ÓO]VEIS[^\n]{0,100})'],t)
        add(fields,'cartorio','Cartório',cart,source)
        cnm=first_match([r'CNM\s*[:\-]?\s*([0-9.\-]{8,})'],t); add(fields,'cnm','CNM',cnm,source)
        inscr=first_match([r'INSCRI[CÇ][AÃ]O IMOBILI[ÁA]RIA\s*[:\-]?\s*([0-9.\-]+)',r'IM[ÓO]VEL CADASTRADO[^\n]{0,80}?N[ºO°.]?\s*([0-9.\-]+)'],t)
        add(fields,'inscricao','Inscrição imobiliária',inscr,source)
        endereco=first_match([r'(?:IM[ÓO]VEL\s*[-–:]\s*)?([^\n]{0,50}(?:RUA|AVENIDA|AV\.)[^\n]{8,150})'],t)
        add(fields,'descricao','Descrição/endereço identificado',endereco,source)
        tipo = 'Apartamento' if re.search(r'\bAPARTAMENTO\b|\bAPTO\.?\b', t, re.I) else ('Casa' if re.search(r'\bCASA RESIDENCIAL\b|\bUMA CASA\b', t, re.I) else ('Terreno' if re.search(r'\bTERRENO\b', t, re.I) else None))
        add(fields,'tipo_imovel','Tipo do imóvel',tipo,source)
        cond=first_match([r'(CONDOM[IÍ]NIO[^\n,;.]{3,100})'],t)
        add(fields,'condominio','Condomínio',cond,source)
        unidade=first_match([r'((?:APARTAMENTO|APTO\.?)\s*(?:N[ºO°.]?\s*)?[0-9A-Z.-]+[^\n]{0,50})'],t)
        add(fields,'complemento','Unidade/Complemento',unidade,source)
        cidadeuf=first_match([r'\b([A-ZÁÀÂÃÉÊÍÓÔÕÚÇ][A-ZÁÀÂÃÉÊÍÓÔÕÚÇ ]{2,})[/-]([A-Z]{2})\b'],t)
        add(fields,'cidade_uf','Cidade/UF',cidadeuf,source)
        owner=first_match([r'PROPRIET[ÁA]RI[OA]\(S\)?\s*[-–:]\s*([^\n]{4,150})',r'PROPRIET[ÁA]RIA\s*[:\-]\s*([^\n]{4,150})'],t)
        add(fields,'proprietario','Proprietário/titular aparente',owner,source)
        # Status cues, never state definitive legal conclusion
        cues=[]
        for label,pat in [('Alienação fiduciária',r'ALIENA[CÇ][AÃ]O FIDUCI[ÁA]RIA'),('Consolidação',r'CONSOLIDA[CÇ][AÃ]O DA PROPRIEDADE'),('Penhora',r'PENHORA'),('Indisponibilidade',r'INDISPONIBILIDADE'),('Cancelamento',r'CANCELAMENTO')]:
            if re.search(pat,t,re.I): cues.append(label)
        if cues: add(fields,'atos','Atos relevantes encontrados',', '.join(cues),source); pending.append('Atos registrais encontrados: confirmar a situação atual antes de usar no contrato.')
        if typ=='Matrícula': pending.append('Confirmar o titular registral mais recente; o sistema não usa automaticamente o primeiro proprietário citado.')
    return fields, list(dict.fromkeys(pending))

@app.post('/api/extract')
async def extract(file: UploadFile = File(...), declared_type: str = Form('auto'), role: str = Form('Outro')):
    data=await file.read(); name=file.filename or 'documento'
    suffix=Path(name).suffix.lower(); text=''; pages=1; method='texto'; ocr_used=False
    try:
        if suffix=='.pdf':
            text,pages,doc=pdf_text(data)
            meaningful=len(re.sub(r'\s+','',text))
            if meaningful < 180 and tesseract_available():
                ocr=ocr_doc(doc)
                if len(ocr)>len(text): text=ocr; method='OCR'; ocr_used=True
        elif suffix in ('.png','.jpg','.jpeg','.webp'):
            text=image_ocr(data); method='OCR' if text else 'imagem'; ocr_used=bool(text)
        else: return JSONResponse({'error':'Formato não suportado'}, status_code=400)
    except Exception as e:
        return JSONResponse({'error':f'Falha ao processar: {e}'}, status_code=400)
    typ=classify(name,text,declared_type)
    fields,pending=extract_fields(typ,text,method)
    # Alguns PDFs digitais trazem apenas o texto do QR/certificado, enquanto os dados pessoais
    # continuam rasterizados na página. Se a primeira leitura não encontrar campos úteis, roda OCR.
    expected_types={'CNH','RG','CIN','CPF','Certidão de casamento','Certidão de nascimento','Comprovante de endereço','Matrícula','BCI/IPTU'}
    useful_keys={f.get('key') for f in fields}
    weak = (typ in expected_types and len(fields) < 2) or (typ in {'CNH','RG','CIN','CPF'} and not ({'nome','cpf'} & useful_keys))
    if suffix=='.pdf' and weak and tesseract_available():
        try:
            if 'doc' not in locals():
                _,_,doc=pdf_text(data)
            ocr=ocr_doc(doc)
            if ocr.strip():
                merged=(text+'\n'+ocr).strip()
                new_fields,new_pending=extract_fields(typ,merged,'OCR')
                if len(new_fields) >= len(fields):
                    text=merged; fields=new_fields; pending=new_pending; method='OCR'; ocr_used=True
        except Exception:
            pass
    if not text.strip(): pending.insert(0,'Não foi possível extrair texto. Instale/ative o Tesseract OCR para documentos escaneados.')
    confidence='Alta' if len(fields)>=4 else ('Média' if len(fields)>=2 else 'Baixa')
    return {'filename':name,'document_type':typ,'role':role,'pages':pages,'method':method,'ocr_used':ocr_used,'confidence':confidence,'fields':fields,'pending':pending,'text_preview':text[:1500]}

class GenerateRequest(BaseModel):
    title: str='CONTRATO PARTICULAR DE PROMESSA DE COMPRA E VENDA DE IMÓVEL'
    text: str
    filename: str='APM_CONTRATO'

@app.post('/api/generate-docx')
def generate_docx(req: GenerateRequest):
    out=Path(tempfile.gettempdir())/(re.sub(r'[^A-Za-z0-9_-]+','_',req.filename)+'.docx')
    d=Document(); sec=d.sections[0]; sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.5);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
    styles=d.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(11)
    lines=req.text.splitlines();
    for i,line in enumerate(lines):
        p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        r=p.add_run(line); r.font.name='Arial'; r.font.size=Pt(11)
        if i==0 or line.strip().startswith('CLÁUSULA') or line.strip() in ('ASSINATURAS','TESTEMUNHAS'):
            r.bold=True
            if i==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.save(out)
    return FileResponse(out,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',filename=out.name)

@app.post('/api/generate-pdf')
def generate_pdf(req: GenerateRequest):
    out=Path(tempfile.gettempdir())/(re.sub(r'[^A-Za-z0-9_-]+','_',req.filename)+'.pdf')
    doc=SimpleDocTemplate(str(out),pagesize=A4,rightMargin=2.5*cm,leftMargin=2.5*cm,topMargin=2.5*cm,bottomMargin=2.5*cm)
    styles=getSampleStyleSheet(); body=ParagraphStyle('body',parent=styles['BodyText'],fontName='Helvetica',fontSize=10.5,leading=15,alignment=TA_JUSTIFY,spaceAfter=7); head=ParagraphStyle('head',parent=body,fontName='Helvetica-Bold',alignment=TA_CENTER,fontSize=12); clause=ParagraphStyle('clause',parent=body,fontName='Helvetica-Bold')
    story=[]
    for i,line in enumerate(req.text.splitlines()):
        esc=(line.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'))
        if not line.strip(): story.append(Spacer(1,4)); continue
        st=head if i==0 else (clause if line.strip().startswith('CLÁUSULA') or line.strip() in ('ASSINATURAS','TESTEMUNHAS') else body)
        story.append(Paragraph(esc,st))
    doc.build(story)
    return FileResponse(out,media_type='application/pdf',filename=out.name)




# ===== Web v0.5: autenticação e persistência central =====
DATABASE_URL=os.getenv('DATABASE_URL', f"sqlite:///{ROOT/'apm_web.db'}")
if DATABASE_URL.startswith('postgres://'):
    DATABASE_URL='postgresql+psycopg2://'+DATABASE_URL[len('postgres://'):]
engine=create_engine(DATABASE_URL, pool_pre_ping=True, future=True, connect_args={'check_same_thread':False} if DATABASE_URL.startswith('sqlite') else {})
SessionLocal=sessionmaker(bind=engine, autoflush=False, autocommit=False)
Base=declarative_base()

class User(Base):
    __tablename__='users'
    id=Column(Integer, primary_key=True)
    email=Column(String(240), unique=True, nullable=False, index=True)
    name=Column(String(240), nullable=False, default='Usuário APM')
    role=Column(String(50), nullable=False, default='corretor')
    password_hash=Column(String(500), nullable=False)
    created_at=Column(DateTime, default=datetime.utcnow)

class Negotiation(Base):
    __tablename__='negotiations'
    id=Column(Integer, primary_key=True)
    code=Column(String(40), unique=True, nullable=False, index=True)
    owner_id=Column(Integer, ForeignKey('users.id'), nullable=False)
    title=Column(String(300), default='')
    buyer=Column(String(300), default='')
    seller=Column(String(300), default='')
    property_name=Column(String(500), default='')
    value_cents=Column(Integer, default=0)
    status=Column(String(80), default='Em elaboração')
    data_json=Column(Text, nullable=False, default='{}')
    created_at=Column(DateTime, default=datetime.utcnow)
    updated_at=Column(DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

Base.metadata.create_all(engine)

def _hash_password(password:str)->str:
    salt=secrets.token_bytes(16)
    dk=hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), salt, 220000)
    return f"pbkdf2_sha256$220000${salt.hex()}${dk.hex()}"

def _verify_password(password:str, stored:str)->bool:
    try:
        alg,it,salt_hex,digest_hex=stored.split('$',3)
        calc=hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), bytes.fromhex(salt_hex), int(it)).hex()
        return hmac.compare_digest(calc,digest_hex)
    except Exception:
        return False

def _seed_admin():
    db=SessionLocal()
    try:
        email=os.getenv('ADMIN_EMAIL','admin@apm.local').strip().lower()
        if not db.query(User).filter(User.email==email).first():
            db.add(User(email=email,name=os.getenv('ADMIN_NAME','Administrador APM'),role='admin',password_hash=_hash_password(os.getenv('ADMIN_PASSWORD','123456'))))
            db.commit()
    finally: db.close()
_seed_admin()

def current_user(request:Request, db):
    uid=request.session.get('user_id')
    return db.query(User).filter(User.id==uid).first() if uid else None

@app.post('/api/login')
async def api_login(request:Request):
    body=await request.json(); email=(body.get('email') or '').strip().lower(); password=body.get('password') or ''
    db=SessionLocal()
    try:
        u=db.query(User).filter(User.email==email).first()
        if not u or not _verify_password(password,u.password_hash):
            return JSONResponse({'ok':False,'message':'E-mail ou senha inválidos.'}, status_code=401)
        request.session['user_id']=u.id
        return {'ok':True,'user':{'id':u.id,'name':u.name,'email':u.email,'role':u.role}}
    finally: db.close()

@app.post('/api/logout')
def api_logout(request:Request):
    request.session.clear(); return {'ok':True}

@app.get('/api/me')
def api_me(request:Request):
    db=SessionLocal()
    try:
        u=current_user(request,db)
        if not u: return JSONResponse({'authenticated':False}, status_code=401)
        return {'authenticated':True,'user':{'id':u.id,'name':u.name,'email':u.email,'role':u.role}}
    finally: db.close()

@app.get('/api/negotiations')
def list_negotiations(request:Request):
    db=SessionLocal()
    try:
        u=current_user(request,db)
        if not u: return JSONResponse({'error':'Não autenticado'}, status_code=401)
        q=db.query(Negotiation)
        if u.role!='admin': q=q.filter(Negotiation.owner_id==u.id)
        rows=q.order_by(Negotiation.updated_at.desc()).all()
        out=[]
        for n in rows:
            try: data=json.loads(n.data_json or '{}')
            except Exception: data={}
            out.append({'id':n.id,'codigo':n.code,'imovel':n.property_name,'comprador':n.buyer,'vendedor':n.seller,'valor':n.value_cents/100,'status':n.status,'data':n.updated_at.strftime('%d/%m/%Y') if n.updated_at else '', 'draft':data})
        return out
    finally: db.close()

@app.post('/api/negotiations')
async def save_negotiation(request:Request):
    db=SessionLocal()
    try:
        u=current_user(request,db)
        if not u: return JSONResponse({'error':'Não autenticado'}, status_code=401)
        body=await request.json(); draft=body.get('draft') or {}; code=(body.get('codigo') or '').strip()
        if not code:
            last=db.query(Negotiation).order_by(Negotiation.id.desc()).first(); code=f"NEG-{(last.id+1 if last else 1):04d}"
        n=db.query(Negotiation).filter(Negotiation.code==code).first()
        if n and u.role!='admin' and n.owner_id!=u.id: return JSONResponse({'error':'Sem permissão'},status_code=403)
        im=draft.get('imovel') or {}; comp=draft.get('comprador') or {}; vend=draft.get('vendedor') or {}
        prop=((im.get('condominio') or im.get('tipo') or 'Imóvel')+' '+(im.get('complemento') or '')).strip()
        if not n:
            n=Negotiation(code=code,owner_id=u.id); db.add(n)
        n.title=body.get('titulo') or prop; n.buyer=comp.get('nome') or 'Comprador'; n.seller=vend.get('nome') or 'Vendedor'; n.property_name=prop
        n.value_cents=int(round(float(draft.get('valor') or 0)*100)); n.status=body.get('status') or 'Em elaboração'; n.data_json=json.dumps(draft,ensure_ascii=False); n.updated_at=datetime.utcnow()
        db.commit(); db.refresh(n)
        return {'ok':True,'id':n.id,'codigo':n.code}
    finally: db.close()

@app.get('/api/negotiations/{neg_id}')
def get_negotiation(neg_id:int, request:Request):
    db=SessionLocal()
    try:
        u=current_user(request,db)
        if not u: return JSONResponse({'error':'Não autenticado'}, status_code=401)
        n=db.query(Negotiation).filter(Negotiation.id==neg_id).first()
        if not n or (u.role!='admin' and n.owner_id!=u.id): return JSONResponse({'error':'Não encontrado'},status_code=404)
        return {'id':n.id,'codigo':n.code,'status':n.status,'draft':json.loads(n.data_json or '{}')}
    finally: db.close()

@app.get('/api/users')
def list_users(request:Request):
    db=SessionLocal()
    try:
        u=current_user(request,db)
        if not u or u.role!='admin': return JSONResponse({'error':'Sem permissão'},status_code=403)
        return [{'id':x.id,'name':x.name,'email':x.email,'role':x.role} for x in db.query(User).order_by(User.name).all()]
    finally: db.close()

@app.post('/api/users')
async def create_user(request:Request):
    db=SessionLocal()
    try:
        u=current_user(request,db)
        if not u or u.role!='admin': return JSONResponse({'error':'Sem permissão'},status_code=403)
        body=await request.json(); email=(body.get('email') or '').strip().lower(); password=body.get('password') or ''
        if not email or len(password)<6: return JSONResponse({'error':'Informe e-mail e senha com pelo menos 6 caracteres.'},status_code=400)
        if db.query(User).filter(User.email==email).first(): return JSONResponse({'error':'E-mail já cadastrado.'},status_code=409)
        nu=User(email=email,name=body.get('name') or email,role=body.get('role') or 'corretor',password_hash=_hash_password(password)); db.add(nu); db.commit(); db.refresh(nu)
        return {'ok':True,'id':nu.id}
    finally: db.close()

# Frontend production routes. Explicit files avoid relative-path/static mount issues on cloud hosts.
app.mount('/assets', StaticFiles(directory=str(ROOT / 'assets')), name='assets')

@app.get('/')
def frontend_index():
    return FileResponse(ROOT / 'index.html', media_type='text/html; charset=utf-8')

@app.get('/styles.css')
def frontend_css():
    return FileResponse(ROOT / 'styles.css', media_type='text/css; charset=utf-8', headers={'Cache-Control':'no-cache'})

@app.get('/app.js')
def frontend_js():
    return FileResponse(ROOT / 'app.js', media_type='application/javascript; charset=utf-8', headers={'Cache-Control':'no-cache'})

@app.get('/favicon.ico', include_in_schema=False)
def favicon():
    logo=ROOT / 'assets' / 'logo.png'
    return FileResponse(logo) if logo.exists() else JSONResponse({}, status_code=204)

